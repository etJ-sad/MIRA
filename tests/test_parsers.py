from docx import Document
from openpyxl import Workbook

from mira.chunking import split_text
from mira.parsers import parse_file


def test_csv_rows_become_labelled_sections(tmp_path):
    path = tmp_path / "stoercodes.csv"
    path.write_text("Code;Beschreibung;Massnahme\nE-417;Übertemperatur Spindel;Lüfter prüfen\n", encoding="utf-8")
    [section] = parse_file(path)
    assert section.text == "Code: E-417; Beschreibung: Übertemperatur Spindel; Massnahme: Lüfter prüfen"
    assert section.location == "Zeile 2"


def test_xlsx_keeps_sheet_and_row(tmp_path):
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Fehler"
    sheet.append(["Code", "Text"])
    sheet.append(["E-1", "Not-Aus aktiv"])
    path = tmp_path / "liste.xlsx"
    workbook.save(path)
    [section] = parse_file(path)
    assert section.location == "Blatt Fehler, Zeile 2"
    assert "Not-Aus aktiv" in section.text


def test_docx_groups_by_heading(tmp_path):
    document = Document()
    document.add_heading("Wartung Pumpe", level=1)
    document.add_paragraph("Dichtung alle 500 h tauschen.")
    path = tmp_path / "pumpe.docx"
    document.save(path)
    [section] = parse_file(path)
    assert section.location == "Wartung Pumpe"
    assert "500 h" in section.text


def test_split_text_overlaps_and_covers_everything():
    text = " ".join(f"Satz {i} über die Anlage." for i in range(200))
    chunks = split_text(text, 300, 50)
    assert len(chunks) > 1
    assert all(len(c) <= 300 for c in chunks)
    assert "Satz 199" in chunks[-1]


def test_blank_rows_keep_original_row_numbers(tmp_path):
    path = tmp_path / "codes.csv"
    path.write_text("Code;Text\n;\nE-9;Lager laut\n", encoding="utf-8")
    [section] = parse_file(path)
    assert section.location == "Zeile 3"

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "F"
    sheet.append(["Code", "Text"])
    sheet.append([None, None])
    sheet.append(["E-9", "Lager laut"])
    xlsx = tmp_path / "codes.xlsx"
    workbook.save(xlsx)
    [section] = parse_file(xlsx)
    assert section.location == "Blatt F, Zeile 3"
