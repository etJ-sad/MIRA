"""Turn files into text sections that remember where they came from."""

from __future__ import annotations

import csv
import io
from dataclasses import dataclass
from pathlib import Path


@dataclass
class Section:
    text: str
    # Human readable position, e.g. "Seite 12" or "Blatt Fehler, Zeile 4"
    location: str


SUPPORTED = {".pdf", ".docx", ".xlsx", ".xlsm", ".csv", ".txt", ".md"}


def parse_file(path: Path) -> list[Section]:
    suffix = path.suffix.lower()
    if suffix == ".pdf":
        return _parse_pdf(path)
    if suffix == ".docx":
        return _parse_docx(path)
    if suffix in {".xlsx", ".xlsm"}:
        return _parse_xlsx(path)
    if suffix == ".csv":
        return _parse_csv(path)
    if suffix in {".txt", ".md"}:
        text = path.read_text(encoding="utf-8", errors="replace")
        return [Section(text, "")] if text.strip() else []
    raise ValueError(f"Dateityp nicht unterstützt: {suffix}")


def _parse_pdf(path: Path) -> list[Section]:
    from pypdf import PdfReader

    reader = PdfReader(str(path))
    sections = []
    for number, page in enumerate(reader.pages, start=1):
        text = page.extract_text() or ""
        if text.strip():
            sections.append(Section(text, f"Seite {number}"))
    return sections


def _parse_docx(path: Path) -> list[Section]:
    import docx

    document = docx.Document(str(path))
    sections: list[Section] = []
    heading = ""
    buffer: list[str] = []

    def flush() -> None:
        if buffer:
            sections.append(Section("\n".join(buffer), heading))
            buffer.clear()

    for paragraph in document.paragraphs:
        text = paragraph.text.strip()
        if not text:
            continue
        if paragraph.style is not None and paragraph.style.name.lower().startswith(("heading", "überschrift")):
            flush()
            heading = text
        buffer.append(text)
    flush()

    for number, table in enumerate(document.tables, start=1):
        rows = [" | ".join(cell.text.strip() for cell in row.cells) for row in table.rows]
        if rows:
            sections.append(Section("\n".join(rows), f"Tabelle {number}"))
    return sections


def _numbered_rows(rows) -> list[tuple[int, list[str]]]:
    """Non-empty rows with their original 1-based row number, so citations match the file."""
    return [(number, row) for number, row in enumerate(rows, start=1) if any(str(cell).strip() for cell in row)]


def _rows_to_sections(header: list[str], rows: list[tuple[int, list[str]]], prefix: str) -> list[Section]:
    """One section per row, written as "Spalte: Wert" so fault-code lists stay searchable."""
    sections = []
    for number, row in rows:
        pairs = [f"{h}: {v}" for h, v in zip(header, row) if str(v).strip()]
        if pairs:
            sections.append(Section("; ".join(pairs), f"{prefix}Zeile {number}"))
    return sections


def _parse_xlsx(path: Path) -> list[Section]:
    from openpyxl import load_workbook

    workbook = load_workbook(str(path), read_only=True, data_only=True)
    sections: list[Section] = []
    for sheet in workbook.worksheets:
        rows = _numbered_rows(["" if v is None else str(v) for v in row] for row in sheet.iter_rows(values_only=True))
        if rows:
            sections += _rows_to_sections(rows[0][1], rows[1:], f"Blatt {sheet.title}, ")
    workbook.close()
    return sections


def _parse_csv(path: Path) -> list[Section]:
    raw = path.read_bytes()
    try:
        text = raw.decode("utf-8-sig")
    except UnicodeDecodeError:
        text = raw.decode("cp1252")
    try:
        dialect = csv.Sniffer().sniff(text[:4096], delimiters=";,\t")
    except csv.Error:
        dialect = csv.excel
    rows = _numbered_rows(csv.reader(io.StringIO(text), dialect))
    if not rows:
        return []
    return _rows_to_sections(rows[0][1], rows[1:], "")
